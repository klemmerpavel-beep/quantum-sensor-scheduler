"""Сверка data/seed_ymg_stage1.json с source/План-график_этап1_ЯМГ-ИИМ.xlsx.

Только чтение. Выводит расхождения по номерам, наименованиям, исполнителям,
срокам, документам, примечаниям и отметкам вкладки «Важное».
"""
import datetime as dt
import json
import re
import sys
import warnings
from pathlib import Path

import openpyxl

warnings.filterwarnings("ignore")
ROOT = Path(__file__).resolve().parent.parent
SEED = json.loads((ROOT / "data/seed_ymg_stage1.json").read_text(encoding="utf-8"))
WB = openpyxl.load_workbook(next((ROOT / "source").glob("*.xlsx")), data_only=True)


def norm_num(v):
    return str(v).strip().rstrip(".").strip() if v is not None else None


def norm_txt(v):
    if v is None:
        return None
    s = re.sub(r"\s+", " ", str(v)).strip()
    return s or None


def to_iso(v):
    """Дата ячейки → (ISO, квалификатор) либо (None, None) для «—»/пусто."""
    if v is None:
        return None, None
    if isinstance(v, dt.datetime):
        return v.date().isoformat(), None
    s = norm_txt(v)
    if s in ("—", "-", "–"):
        return None, None
    m = re.search(r"(\d{2})\.(\d{2})\.(\d{4})", s)
    if not m:
        return "?" + s, None
    iso = f"{m.group(3)}-{m.group(2)}-{m.group(1)}"
    q = None
    if "условно" in s:
        q = "условно"
    elif "не позднее" in s:
        q = "не позднее"
    return iso, q


items = {i["num"]: i for i in SEED["items"]}
diffs = []

# ── Вкладка «План-график (На подписание)»
ws = WB["План-график (На подписание)"]
plan_nums = []
for r in ws.iter_rows(min_row=10, max_row=81):
    num = norm_num(r[0].value)
    if not num:
        continue
    plan_nums.append(num)
    it = items.get(num)
    if it is None:
        diffs.append(("ПГ", num, "нет в JSON", None, None))
        continue
    name = norm_txt(r[1].value)
    if name != norm_txt(it["name"]):
        diffs.append(("ПГ", num, "name", name, it["name"]))
    if it["kind"] == "section":
        continue
    owner = norm_txt(r[2].value)
    if owner != norm_txt(it["owner_raw"]):
        diffs.append(("ПГ", num, "owner_raw", owner, it["owner_raw"]))
    due, dq = to_iso(r[3].value)
    if due != it["due"] or dq != it["due_qualifier"]:
        diffs.append(("ПГ", num, "due", (due, dq), (it["due"], it["due_qualifier"])))
    for col, key in ((4, "output_doc"), (6, "execution_doc"), (7, "basis"), (8, "plan_note")):
        a = norm_txt(r[col].value)
        a = None if a in ("—",) else a
        if a != norm_txt(it[key]):
            diffs.append(("ПГ", num, key, a, it[key]))
    dtc, dtq = to_iso(r[5].value)
    if dtc != it["due_to_customer"] or dtq != it["due_to_customer_qualifier"]:
        diffs.append(("ПГ", num, "due_to_customer", (dtc, dtq), (it["due_to_customer"], it["due_to_customer_qualifier"])))

missing = [n for n in items if n not in plan_nums]
order_ok = plan_nums == [i["num"] for i in SEED["items"]]

# ── Вкладка «Важное»
wv = WB["Важное"]
block = None
vazhnoe = {}
for r in wv.iter_rows(min_row=4, max_row=wv.max_row):
    a = norm_txt(r[0].value)
    if a and not re.match(r"^\d", a) and r[1].value is None:
        block = a.split("—")[0].strip()
        continue
    if a and a.startswith("№"):
        continue
    if not any(c.value for c in r):
        continue
    num = norm_num(r[0].value)
    due, _ = to_iso(r[3].value)
    tr, _ = to_iso(r[5].value)
    vazhnoe[num or ("вне плана: " + norm_txt(r[1].value)[:40])] = dict(
        block=block, due=due, transfer=tr, mark=norm_txt(r[6].value), owner=norm_txt(r[2].value))

vdiffs = []
for num, v in vazhnoe.items():
    it = items.get(num)
    if it is None:
        vdiffs.append((num, "вне План-графика", v))
        continue
    if v["due"] and v["due"] != it["due"]:
        vdiffs.append((num, "срок «Важного» ≠ срок ПГ", (v["due"], it["due"])))
    mark = v["mark"]
    if (mark or None) != (norm_txt(it["status_mark"]) or None):
        vdiffs.append((num, "status_mark", (mark, it["status_mark"])))
    t = v["transfer"]
    stored = it["closed_date"] if it["status"] == "Закрыто" else it["control_date"]
    if t != stored:
        vdiffs.append((num, "Дата передачи", (t, stored, it["status"])))
not_in_vazhnoe = [n for n, i in items.items() if i["kind"] != "section" and n not in vazhnoe]

out = dict(
    plan_rows=len(plan_nums), json_items=len(items), order_identical=order_ok,
    json_missing_in_plan=missing, plan_diffs=diffs,
    vazhnoe_rows=len(vazhnoe), vazhnoe_diffs=vdiffs, not_in_vazhnoe=not_in_vazhnoe,
)
json.dump(out, sys.stdout, ensure_ascii=False, indent=1, default=str)
